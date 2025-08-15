"""
Document Comparison Service for MinuteMate
Compares original and edited documents to learn user preferences
"""

import logging
import json
import re
from datetime import datetime, timezone
from typing import Dict, List, Any, Optional, Tuple
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import difflib
from models import db, UserPreference, Meeting

logger = logging.getLogger(__name__)

class DocumentComparisonService:
    """Service for comparing documents and learning user preferences"""
    
    def __init__(self):
        self.preference_categories = {
            'formatting': ['font_family', 'font_size', 'line_spacing', 'paragraph_spacing'],
            'structure': ['heading_style', 'section_order', 'bullet_style', 'numbering_style'],
            'content': ['language_style', 'formality_level', 'detail_level', 'action_item_format'],
            'layout': ['margins', 'page_orientation', 'header_footer', 'table_style']
        }
    
    def compare_documents(self, original_path: str, edited_path: str, user_id: str, meeting_id: str) -> Dict[str, Any]:
        """
        Compare original and edited documents to extract user preferences
        
        Args:
            original_path: Path to original document
            edited_path: Path to edited document
            user_id: User ID for storing preferences
            meeting_id: Meeting ID for reference
            
        Returns:
            Dictionary containing comparison results and learned preferences
        """
        try:
            logger.info(f"Starting document comparison for user {user_id}, meeting {meeting_id}")
            
            # Load documents
            original_doc = Document(original_path)
            edited_doc = Document(edited_path)
            
            # Extract document properties
            original_props = self._extract_document_properties(original_doc)
            edited_props = self._extract_document_properties(edited_doc)
            
            # Compare content structure
            content_changes = self._compare_content_structure(original_doc, edited_doc)
            
            # Compare formatting
            formatting_changes = self._compare_formatting(original_doc, edited_doc)
            
            # Analyze text changes
            text_changes = self._analyze_text_changes(original_doc, edited_doc)
            
            # Extract preferences
            learned_preferences = self._extract_preferences(
                original_props, edited_props, content_changes, formatting_changes, text_changes
            )
            
            # Store preferences
            self._store_user_preferences(user_id, learned_preferences)
            
            # Create comparison report
            comparison_result = {
                'meeting_id': meeting_id,
                'user_id': user_id,
                'comparison_date': datetime.now(timezone.utc).isoformat(),
                'original_properties': original_props,
                'edited_properties': edited_props,
                'content_changes': content_changes,
                'formatting_changes': formatting_changes,
                'text_changes': text_changes,
                'learned_preferences': learned_preferences,
                'preference_confidence': self._calculate_confidence(learned_preferences)
            }
            
            logger.info(f"Document comparison completed. Learned {len(learned_preferences)} preferences")
            return comparison_result
            
        except Exception as e:
            logger.error(f"Document comparison failed: {str(e)}")
            raise
    
    def _extract_document_properties(self, doc: Document) -> Dict[str, Any]:
        """Extract properties from a document"""
        properties = {
            'paragraph_count': len(doc.paragraphs),
            'table_count': len(doc.tables),
            'styles_used': [],
            'font_families': set(),
            'font_sizes': set(),
            'alignments': set(),
            'heading_levels': set()
        }
        
        for paragraph in doc.paragraphs:
            if paragraph.style:
                properties['styles_used'].append(paragraph.style.name)
                
            if paragraph.style.name.startswith('Heading'):
                level = re.findall(r'\d+', paragraph.style.name)
                if level:
                    properties['heading_levels'].add(int(level[0]))
            
            for run in paragraph.runs:
                if run.font.name:
                    properties['font_families'].add(run.font.name)
                if run.font.size:
                    properties['font_sizes'].add(run.font.size.pt)
            
            if paragraph.alignment:
                properties['alignments'].add(str(paragraph.alignment))
        
        # Convert sets to lists for JSON serialization
        properties['font_families'] = list(properties['font_families'])
        properties['font_sizes'] = list(properties['font_sizes'])
        properties['alignments'] = list(properties['alignments'])
        properties['heading_levels'] = list(properties['heading_levels'])
        properties['styles_used'] = list(set(properties['styles_used']))
        
        return properties
    
    def _compare_content_structure(self, original_doc: Document, edited_doc: Document) -> Dict[str, Any]:
        """Compare the content structure between documents"""
        original_structure = self._extract_structure(original_doc)
        edited_structure = self._extract_structure(edited_doc)
        
        changes = {
            'sections_added': [],
            'sections_removed': [],
            'sections_reordered': False,
            'heading_changes': [],
            'list_format_changes': []
        }
        
        # Compare sections
        original_sections = set(original_structure['sections'])
        edited_sections = set(edited_structure['sections'])
        
        changes['sections_added'] = list(edited_sections - original_sections)
        changes['sections_removed'] = list(original_sections - edited_sections)
        
        # Check if sections were reordered
        common_sections = original_sections & edited_sections
        if len(common_sections) > 1:
            original_order = [s for s in original_structure['sections'] if s in common_sections]
            edited_order = [s for s in edited_structure['sections'] if s in common_sections]
            changes['sections_reordered'] = original_order != edited_order
        
        return changes
    
    def _extract_structure(self, doc: Document) -> Dict[str, Any]:
        """Extract structural elements from document"""
        structure = {
            'sections': [],
            'headings': [],
            'lists': [],
            'tables': []
        }
        
        for paragraph in doc.paragraphs:
            if paragraph.style.name.startswith('Heading'):
                structure['headings'].append({
                    'text': paragraph.text,
                    'level': paragraph.style.name,
                    'alignment': str(paragraph.alignment) if paragraph.alignment else None
                })
                structure['sections'].append(paragraph.text.strip())
            elif paragraph.style.name.startswith('List'):
                structure['lists'].append({
                    'text': paragraph.text,
                    'style': paragraph.style.name
                })
        
        for table in doc.tables:
            structure['tables'].append({
                'rows': len(table.rows),
                'cols': len(table.columns)
            })
        
        return structure
    
    def _compare_formatting(self, original_doc: Document, edited_doc: Document) -> Dict[str, Any]:
        """Compare formatting between documents"""
        original_formatting = self._extract_formatting(original_doc)
        edited_formatting = self._extract_formatting(edited_doc)
        
        changes = {}
        
        for category in ['fonts', 'sizes', 'alignments', 'spacing']:
            if category in original_formatting and category in edited_formatting:
                changes[f'{category}_changes'] = {
                    'original': original_formatting[category],
                    'edited': edited_formatting[category],
                    'changed': original_formatting[category] != edited_formatting[category]
                }
        
        return changes
    
    def _extract_formatting(self, doc: Document) -> Dict[str, Any]:
        """Extract formatting information from document"""
        formatting = {
            'fonts': {},
            'sizes': {},
            'alignments': {},
            'spacing': {}
        }
        
        for paragraph in doc.paragraphs:
            # Track font usage
            for run in paragraph.runs:
                if run.font.name:
                    font_name = run.font.name
                    formatting['fonts'][font_name] = formatting['fonts'].get(font_name, 0) + 1
                
                if run.font.size:
                    size = run.font.size.pt
                    formatting['sizes'][size] = formatting['sizes'].get(size, 0) + 1
            
            # Track alignment
            if paragraph.alignment:
                align = str(paragraph.alignment)
                formatting['alignments'][align] = formatting['alignments'].get(align, 0) + 1
        
        return formatting
    
    def _analyze_text_changes(self, original_doc: Document, edited_doc: Document) -> Dict[str, Any]:
        """Analyze text-level changes between documents"""
        original_text = '\n'.join([p.text for p in original_doc.paragraphs])
        edited_text = '\n'.join([p.text for p in edited_doc.paragraphs])
        
        # Use difflib to find differences
        differ = difflib.unified_diff(
            original_text.splitlines(keepends=True),
            edited_text.splitlines(keepends=True),
            fromfile='original',
            tofile='edited'
        )
        
        changes = {
            'additions': [],
            'deletions': [],
            'modifications': [],
            'word_count_change': len(edited_text.split()) - len(original_text.split()),
            'character_count_change': len(edited_text) - len(original_text)
        }
        
        diff_lines = list(differ)
        for line in diff_lines:
            if line.startswith('+') and not line.startswith('+++'):
                changes['additions'].append(line[1:].strip())
            elif line.startswith('-') and not line.startswith('---'):
                changes['deletions'].append(line[1:].strip())
        
        return changes
    
    def _extract_preferences(self, original_props: Dict, edited_props: Dict, 
                           content_changes: Dict, formatting_changes: Dict, 
                           text_changes: Dict) -> Dict[str, Any]:
        """Extract user preferences from comparison results"""
        preferences = {}
        
        # Font preferences
        if 'fonts_changes' in formatting_changes and formatting_changes['fonts_changes']['changed']:
            edited_fonts = formatting_changes['fonts_changes']['edited']
            if edited_fonts:
                most_used_font = max(edited_fonts.items(), key=lambda x: x[1])
                preferences['preferred_font'] = most_used_font[0]
        
        # Font size preferences
        if 'sizes_changes' in formatting_changes and formatting_changes['sizes_changes']['changed']:
            edited_sizes = formatting_changes['sizes_changes']['edited']
            if edited_sizes:
                most_used_size = max(edited_sizes.items(), key=lambda x: x[1])
                preferences['preferred_font_size'] = most_used_size[0]
        
        # Alignment preferences
        if 'alignments_changes' in formatting_changes and formatting_changes['alignments_changes']['changed']:
            edited_alignments = formatting_changes['alignments_changes']['edited']
            if edited_alignments:
                most_used_alignment = max(edited_alignments.items(), key=lambda x: x[1])
                preferences['preferred_alignment'] = most_used_alignment[0]
        
        # Content preferences
        if content_changes['sections_added']:
            preferences['preferred_sections'] = content_changes['sections_added']
        
        if content_changes['sections_reordered']:
            preferences['prefers_custom_section_order'] = True
        
        # Text style preferences
        if text_changes['word_count_change'] > 0:
            preferences['prefers_detailed_content'] = True
        elif text_changes['word_count_change'] < -50:  # Significant reduction
            preferences['prefers_concise_content'] = True
        
        # Heading preferences
        if edited_props['heading_levels'] != original_props['heading_levels']:
            preferences['preferred_heading_levels'] = edited_props['heading_levels']
        
        return preferences
    
    def _calculate_confidence(self, preferences: Dict[str, Any]) -> float:
        """Calculate confidence score for learned preferences"""
        if not preferences:
            return 0.0
        
        # Base confidence on number of preferences learned
        base_confidence = min(len(preferences) * 0.1, 0.8)
        
        # Boost confidence for structural changes (more reliable)
        structural_prefs = ['preferred_sections', 'prefers_custom_section_order', 'preferred_heading_levels']
        structural_count = sum(1 for pref in structural_prefs if pref in preferences)
        structural_boost = structural_count * 0.1
        
        return min(base_confidence + structural_boost, 1.0)
    
    def _store_user_preferences(self, user_id: str, preferences: Dict[str, Any]):
        """Store learned preferences in the database"""
        try:
            for pref_key, pref_value in preferences.items():
                # Check if preference already exists
                existing_pref = UserPreference.query.filter_by(
                    user_id=user_id,
                    preference_key=pref_key
                ).first()
                
                if existing_pref:
                    # Update existing preference
                    existing_pref.preference_value = json.dumps(pref_value)
                    existing_pref.confidence_score = min(existing_pref.confidence_score + 0.1, 1.0)
                    existing_pref.updated_at = datetime.now(timezone.utc)
                else:
                    # Create new preference
                    new_pref = UserPreference(
                        user_id=user_id,
                        preference_key=pref_key,
                        preference_value=json.dumps(pref_value),
                        confidence_score=0.5,  # Initial confidence
                        category='learned_from_edit'
                    )
                    db.session.add(new_pref)
            
            db.session.commit()
            logger.info(f"Stored {len(preferences)} preferences for user {user_id}")
            
        except Exception as e:
            db.session.rollback()
            logger.error(f"Failed to store preferences: {str(e)}")
            raise
    
    def get_user_preferences(self, user_id: str) -> Dict[str, Any]:
        """Get all learned preferences for a user"""
        try:
            preferences = UserPreference.query.filter_by(user_id=user_id).all()
            
            result = {}
            for pref in preferences:
                try:
                    result[pref.preference_key] = {
                        'value': json.loads(pref.preference_value),
                        'confidence': pref.confidence_score,
                        'category': pref.category,
                        'updated_at': pref.updated_at.isoformat()
                    }
                except json.JSONDecodeError:
                    # Handle non-JSON values
                    result[pref.preference_key] = {
                        'value': pref.preference_value,
                        'confidence': pref.confidence_score,
                        'category': pref.category,
                        'updated_at': pref.updated_at.isoformat()
                    }
            
            return result
            
        except Exception as e:
            logger.error(f"Failed to get user preferences: {str(e)}")
            return {}
    
    def apply_preferences_to_document(self, doc_path: str, user_id: str) -> str:
        """Apply learned user preferences to a document"""
        try:
            preferences = self.get_user_preferences(user_id)
            
            if not preferences:
                logger.info(f"No preferences found for user {user_id}")
                return doc_path
            
            doc = Document(doc_path)
            
            # Apply font preferences
            if 'preferred_font' in preferences:
                font_name = preferences['preferred_font']['value']
                self._apply_font_preference(doc, font_name)
            
            # Apply font size preferences
            if 'preferred_font_size' in preferences:
                font_size = preferences['preferred_font_size']['value']
                self._apply_font_size_preference(doc, font_size)
            
            # Apply alignment preferences
            if 'preferred_alignment' in preferences:
                alignment = preferences['preferred_alignment']['value']
                self._apply_alignment_preference(doc, alignment)
            
            # Save modified document
            modified_path = doc_path.replace('.docx', '_personalized.docx')
            doc.save(modified_path)
            
            logger.info(f"Applied preferences to document: {modified_path}")
            return modified_path
            
        except Exception as e:
            logger.error(f"Failed to apply preferences: {str(e)}")
            return doc_path
    
    def _apply_font_preference(self, doc: Document, font_name: str):
        """Apply font preference to document"""
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if run.font.name != font_name:
                    run.font.name = font_name
    
    def _apply_font_size_preference(self, doc: Document, font_size: float):
        """Apply font size preference to document"""
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if not paragraph.style.name.startswith('Heading'):  # Don't change heading sizes
                    run.font.size = Pt(font_size)
    
    def _apply_alignment_preference(self, doc: Document, alignment: str):
        """Apply alignment preference to document"""
        # Map string alignment to enum
        alignment_map = {
            'WD_ALIGN_PARAGRAPH.LEFT': WD_ALIGN_PARAGRAPH.LEFT,
            'WD_ALIGN_PARAGRAPH.CENTER': WD_ALIGN_PARAGRAPH.CENTER,
            'WD_ALIGN_PARAGRAPH.RIGHT': WD_ALIGN_PARAGRAPH.RIGHT,
            'WD_ALIGN_PARAGRAPH.JUSTIFY': WD_ALIGN_PARAGRAPH.JUSTIFY
        }
        
        if alignment in alignment_map:
            target_alignment = alignment_map[alignment]
            for paragraph in doc.paragraphs:
                if not paragraph.style.name.startswith('Heading'):  # Don't change heading alignment
                    paragraph.alignment = target_alignment
